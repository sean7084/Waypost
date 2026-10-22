"""
Tests for the inspections web UI: batches, export, dashboard, permissions.

Complements ``inspections/tests.py`` (models, transforms, report generator,
import command) by exercising the server-rendered views and the new
``InspectionBatch`` grouping.
"""
import io
from datetime import date, timedelta

import pandas as pd
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import AdminRole
from companies.models import Company, Division, Location
from inspections.forms import AssetListExportForm, InspectionBatchByBrandForm
from inspections.models import InspectionBatch, InspectionDevice, StoreInspection
from inspections.services.asset_list_export import export_asset_list
from tests.storage import IsolatedMediaTestCase

User = get_user_model()

# Role codes used by the tests; ensured to exist in setUpTestData.
_ROLE_CODES = ['superadmin', 'it_administrator', 'inspection_engineer']


def _ensure_roles():
    """Make sure the AdminRole rows the tests assign actually exist."""
    defaults = {
        'superadmin': ('Superadmin', 'Full access.'),
        'it_administrator': ('IT Administrator', 'Scoped admin.'),
        'inspection_engineer': ('Inspection Engineer', 'Onsite engineer.'),
    }
    for code in _ROLE_CODES:
        name, description = defaults[code]
        AdminRole.objects.get_or_create(
            code=code, defaults={'name': name, 'description': description, 'is_active': True}
        )


def _make_user(username, role_code=None, **kwargs):
    user = User.objects.create_user(username=username, password='pw', **kwargs)
    if role_code:
        user.set_admin_roles([role_code])
    return user


def _schedule_xlsx_bytes(rows):
    """Build an in-memory schedule.xlsx and return its bytes."""
    df = pd.DataFrame(rows)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
    return buffer.getvalue()


def _assets_xlsx_bytes(rows):
    df = pd.DataFrame(rows)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
    return buffer.getvalue()


class InspectionFrontendTestBase(IsolatedMediaTestCase):
    """Shared fixtures: company, division, locations, users with roles."""

    @classmethod
    def setUpTestData(cls):
        _ensure_roles()
        cls.company = Company.objects.create(name='Kering', code='KER')
        cls.division = Division.objects.create(company=cls.company, name='Gucci', code='GUCCI')
        # Three active stores + one closed + one office (should be excluded).
        cls.locations = [
            Location.objects.create(
                company=cls.company, division=cls.division, name=f'Store{i}',
                code=f'2200{i}', location_type=Location.LocationType.STORE,
                status=Location.LocationStatus.ACTIVE,
            )
            for i in range(1, 4)
        ]
        Location.objects.create(
            company=cls.company, division=cls.division, name='ClosedStore',
            code='22099', location_type=Location.LocationType.STORE,
            status=Location.LocationStatus.CLOSED,
        )
        Location.objects.create(
            company=cls.company, division=cls.division, name='Office',
            code='22098', location_type=Location.LocationType.OFFICE,
            status=Location.LocationStatus.ACTIVE,
        )

        cls.superadmin = _make_user('admin', 'superadmin')
        cls.it_admin = _make_user('itadmin', 'it_administrator')
        cls.engineer = _make_user('engineer', 'inspection_engineer')
        cls.outsider = _make_user('outsider')  # no inspection role


# -- batch by-brand flow -----------------------------------------------------
class BatchByBrandTests(InspectionFrontendTestBase):
    def test_form_creates_batch_and_spreads_stores(self):
        self.client.force_login(self.superadmin)
        start = date.today()
        end = start + timedelta(days=2)  # 3 days, 3 stores -> one per day
        response = self.client.post(reverse('inspections:batch_create_by_brand'), {
            'name': 'Gucci Test Batch',
            'company': self.company.id,
            'division': self.division.id,
            'start_date': start.isoformat(),
            'end_date': end.isoformat(),
            'description': 'unit test',
        })
        batch = InspectionBatch.objects.get(name='Gucci Test Batch')
        self.assertRedirects(response, reverse('inspections:batch_detail', args=[batch.id]))
        self.assertEqual(batch.source, InspectionBatch.Source.BRAND)
        self.assertEqual(batch.division, self.division)
        # Only the 3 active STORE locations should be scheduled.
        self.assertEqual(batch.inspections.count(), 3)
        # Each store lands on a distinct date (round-robin across 3 days).
        dates = sorted(i.inspection_date for i in batch.inspections.all())
        self.assertEqual(dates, [start, start + timedelta(days=1), start + timedelta(days=2)])

    def test_form_rejects_division_company_mismatch(self):
        other_company = Company.objects.create(name='Other', code='OTH')
        other_division = Division.objects.create(company=other_company, name='OtherBrand', code='OB')
        form = InspectionBatchByBrandForm(data={
            'name': 'Bad', 'company': self.company.id, 'division': other_division.id,
            'start_date': date.today().isoformat(),
            'end_date': (date.today() + timedelta(days=1)).isoformat(),
        }, user=self.it_admin)
        self.assertFalse(form.is_valid())
        self.assertIn('division', form.errors)

    def test_form_rejects_end_before_start(self):
        form = InspectionBatchByBrandForm(data={
            'name': 'Bad', 'company': self.company.id, 'division': self.division.id,
            'start_date': date.today().isoformat(),
            'end_date': (date.today() - timedelta(days=1)).isoformat(),
        }, user=self.it_admin)
        self.assertFalse(form.is_valid())
        self.assertIn('end_date', form.errors)

    def test_engineer_cannot_create_batch(self):
        self.client.force_login(self.engineer)
        response = self.client.get(reverse('inspections:batch_create_by_brand'))
        self.assertEqual(response.status_code, 403)

    def test_outsider_cannot_view_batch_list(self):
        self.client.force_login(self.outsider)
        response = self.client.get(reverse('inspections:batch_list'))
        self.assertEqual(response.status_code, 403)


# -- batch import flow -------------------------------------------------------
class BatchImportTests(InspectionFrontendTestBase):
    def _schedule_rows(self):
        return [
            {'Brand': 'Gucci', 'Store Name': 'Store1', 'JDA code': 22001,
             'inspection_date': date.today()},
            {'Brand': 'Gucci', 'Store Name': 'Store2', 'JDA code': 22002,
             'inspection_date': date.today() + timedelta(days=1)},
        ]

    def test_schedule_only_import_creates_empty_inspections(self):
        self.client.force_login(self.superadmin)
        upload = SimpleUploadedFile(
            'schedule.xlsx', _schedule_xlsx_bytes(self._schedule_rows()),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response = self.client.post(reverse('inspections:batch_import'), {
            'name': 'Import Test', 'company': self.company.id, 'schedule_file': upload,
        })
        batch = InspectionBatch.objects.get(name='Import Test')
        self.assertRedirects(response, reverse('inspections:batch_detail', args=[batch.id]))
        self.assertEqual(batch.source, InspectionBatch.Source.UPLOAD)
        self.assertEqual(batch.inspections.count(), 2)
        # No asset list -> no expected devices.
        self.assertEqual(InspectionDevice.objects.count(), 0)
        # Batch date range tightened to match imported inspections.
        self.assertEqual(batch.start_date, date.today())
        self.assertEqual(batch.end_date, date.today() + timedelta(days=1))

    def test_schedule_plus_assets_import_populates_devices(self):
        self.client.force_login(self.superadmin)
        schedule = SimpleUploadedFile(
            'schedule.xlsx', _schedule_xlsx_bytes(self._schedule_rows()),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        assets = SimpleUploadedFile(
            'assets.xlsx',
            _assets_xlsx_bytes([
                {'STORE': 'Store1', 'JDA': 22001, 'Asset ID': 'A1', 'Category': 'Desktop',
                 'Brand-Model': 'Lenovo-M70Q', 'SN': 'PC1', 'Warranty-start': '',
                 'Usage': 'xstore', 'Status': 'In Store', '照片需求': '需要拍照',
                 '大纲': '', '注释': ''},
            ]),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response = self.client.post(reverse('inspections:batch_import'), {
            'name': 'Full Import', 'company': self.company.id,
            'schedule_file': schedule, 'assets_file': assets,
        })
        batch = InspectionBatch.objects.get(name='Full Import')
        self.assertRedirects(response, reverse('inspections:batch_detail', args=[batch.id]))
        self.assertEqual(InspectionDevice.objects.count(), 1)
        device = InspectionDevice.objects.first()
        self.assertEqual(device.sn, 'PC1')
        self.assertEqual(device.store_inspection.batch, batch)

    def test_engineer_cannot_import(self):
        self.client.force_login(self.engineer)
        response = self.client.get(reverse('inspections:batch_import'))
        self.assertEqual(response.status_code, 403)


# -- asset list export -------------------------------------------------------
class AssetListExportTests(InspectionFrontendTestBase):
    def _make_inspection_with_devices(self, location, inspection_date, batch=None):
        inspection = StoreInspection.objects.create(
            company=self.company, division=self.division, location=location,
            jda_code=location.code, brand_name=self.division.name,
            store_name=location.name, inspection_date=inspection_date,
            batch=batch, status=StoreInspection.Status.PLANNED,
        )
        # One collected device, one untouched device.
        InspectionDevice.objects.create(
            store_inspection=inspection, client_device_uid='c1',
            category='Desktop', brand_model='Lenovo-M70Q', sn='PC1',
            asset_id_text='A1', usage='xstore',
        )
        InspectionDevice.objects.create(
            store_inspection=inspection, client_device_uid='u1',
            category='Monitor', brand_model='Lenovo-L1710', sn='MON1',
            asset_id_text='A2', usage='xstore',
        )
        # Mark the first as collected.
        from django.utils import timezone
        InspectionDevice.objects.filter(client_device_uid='c1').update(collected_at=timezone.now())
        return inspection

    def test_export_includes_untouched_as_not_in_store(self):
        inspection = self._make_inspection_with_devices(
            self.locations[0], date.fromisocalendar(2026, 5, 3)  # Wed of ISO week 5
        )
        content, filename = export_asset_list([inspection])
        self.assertTrue(filename.endswith('.xlsx'))
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook['asset_list']
        # Header row + 2 device rows.
        self.assertEqual(sheet.max_row, 3)
        # Find the Status column.
        header = {sheet.cell(row=1, column=c).value: c for c in range(1, sheet.max_column + 1)}
        status_col = header['Status']
        statuses = sorted(sheet.cell(row=r, column=status_col).value for r in (2, 3))
        self.assertEqual(statuses, ['In Store', 'Not In Store'])

    def test_export_filename_single_brand_single_week(self):
        # All inspections in ISO week 5 of 2026, same division.
        inspections = [
            self._make_inspection_with_devices(self.locations[0], date.fromisocalendar(2026, 5, 1)),
            self._make_inspection_with_devices(self.locations[1], date.fromisocalendar(2026, 5, 2)),
        ]
        _content, filename = export_asset_list(inspections)
        self.assertEqual(filename, 'Gucci_资产表_2026 W5.xlsx')

    def test_export_filename_falls_back_to_batch_name(self):
        batch = InspectionBatch.objects.create(
            name='Multi Week Batch', company=self.company, division=self.division,
            start_date=date(2026, 1, 1), end_date=date(2026, 12, 31),
        )
        inspections = [
            self._make_inspection_with_devices(self.locations[0], date(2026, 1, 5), batch=batch),
            self._make_inspection_with_devices(self.locations[1], date(2026, 6, 15), batch=batch),
        ]
        _content, filename = export_asset_list(inspections, batch=batch)
        self.assertEqual(filename, 'Multi Week Batch.xlsx')

    def test_export_view_requires_at_least_one_filter(self):
        form = AssetListExportForm(data={}, user=self.it_admin)
        self.assertFalse(form.is_valid())
        self.assertIn('__all__', form.errors)

    def test_export_view_week_filter_resolves_inspections(self):
        self._make_inspection_with_devices(self.locations[0], date.fromisocalendar(2026, 5, 3))
        self._make_inspection_with_devices(self.locations[1], date.fromisocalendar(2026, 6, 3))
        form = AssetListExportForm(data={'week': '2026-W05'}, user=self.it_admin)
        self.assertTrue(form.is_valid(), form.errors)
        qs = form.resolve_inspections()
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first().inspection_date, date.fromisocalendar(2026, 5, 3))

    def test_export_view_post_streams_xlsx(self):
        batch = InspectionBatch.objects.create(
            name='Export Batch', company=self.company, division=self.division,
            start_date=date.today(), end_date=date.today() + timedelta(days=1),
        )
        self._make_inspection_with_devices(self.locations[0], date.today(), batch=batch)
        self.client.force_login(self.superadmin)
        response = self.client.post(reverse('inspections:asset_list_export'), {'batch': batch.id})
        self.assertEqual(response.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                      response['Content-Type'])
        # Non-ASCII filenames are RFC 2047 encoded; just confirm a disposition is set.
        self.assertTrue(response['Content-Disposition'])


# -- dashboard + permissions -------------------------------------------------
class DashboardPermissionTests(InspectionFrontendTestBase):
    def test_anonymous_redirected_to_login(self):
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertEqual(response.status_code, 302)
        # LOGIN_URL may be '/login/' or '/accounts/login/'; both are acceptable.
        self.assertIn('login', response['Location'])

    def test_outsider_forbidden(self):
        self.client.force_login(self.outsider)
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertEqual(response.status_code, 403)

    def test_engineer_sees_only_own_inspections(self):
        # One inspection assigned to the engineer, one to someone else.
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), engineer=self.engineer,
            status=StoreInspection.Status.PLANNED,
        )
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[1],
            jda_code='22002', brand_name='Gucci', store_name='Store2',
            inspection_date=date.today(), engineer=self.it_admin,
            status=StoreInspection.Status.PLANNED,
        )
        self.client.force_login(self.engineer)
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertEqual(response.status_code, 200)
        # The engineer's site card should be in the payload; the other should not.
        days_json = response.context['days_json']
        self.assertIn('Store1', days_json)
        self.assertNotIn('22002 Gucci Store2', days_json)

    def test_admin_sees_all_inspections(self):
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), engineer=self.engineer,
            status=StoreInspection.Status.PLANNED,
        )
        self.client.force_login(self.superadmin)
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('Store1', response.context['days_json'])

    def test_dashboard_batch_filter(self):
        batch = InspectionBatch.objects.create(
            name='Filtered Batch', company=self.company, division=self.division,
            start_date=date.today(), end_date=date.today() + timedelta(days=1),
        )
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), batch=batch,
            status=StoreInspection.Status.PLANNED,
        )
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[1],
            jda_code='22002', brand_name='Gucci', store_name='Store2',
            inspection_date=date.today(),
            status=StoreInspection.Status.PLANNED,
        )
        self.client.force_login(self.superadmin)
        response = self.client.get(reverse('inspections:dashboard'), {'batch': batch.id})
        self.assertEqual(response.status_code, 200)
        days_json = response.context['days_json']
        self.assertIn('Store1', days_json)
        self.assertNotIn('Store2', days_json)


# -- inspection list / detail ------------------------------------------------
class InspectionListDetailTests(InspectionFrontendTestBase):
    def test_list_view_renders(self):
        StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), status=StoreInspection.Status.PLANNED,
        )
        self.client.force_login(self.it_admin)
        response = self.client.get(reverse('inspections:inspection_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Store1')

    def test_detail_view_renders_devices(self):
        inspection = StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), status=StoreInspection.Status.PLANNED,
        )
        InspectionDevice.objects.create(
            store_inspection=inspection, client_device_uid='d1',
            category='Desktop', brand_model='Lenovo-M70Q', sn='PC1',
        )
        self.client.force_login(self.it_admin)
        response = self.client.get(reverse('inspections:inspection_detail', args=[inspection.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'PC1')
        self.assertContains(response, 'Lenovo-M70Q')

    def test_engineer_cannot_see_others_inspection_detail(self):
        inspection = StoreInspection.objects.create(
            company=self.company, division=self.division, location=self.locations[0],
            jda_code='22001', brand_name='Gucci', store_name='Store1',
            inspection_date=date.today(), engineer=self.it_admin,
            status=StoreInspection.Status.PLANNED,
        )
        self.client.force_login(self.engineer)
        response = self.client.get(reverse('inspections:inspection_detail', args=[inspection.id]))
        # Engineer's scoped queryset excludes this inspection -> 404.
        self.assertEqual(response.status_code, 404)


# -- navigation gating -------------------------------------------------------
class NavPermissionTests(InspectionFrontendTestBase):
    def test_nav_visible_for_admin(self):
        self.client.force_login(self.it_admin)
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertContains(response, 'Inspections')
        # The nav entry was renamed "Export Asset List" -> "Export", and a new
        # "Review" entry was added for managers.
        self.assertContains(response, reverse('inspections:asset_list_export'))
        self.assertContains(response, reverse('inspections:review'))

    def test_nav_visible_for_engineer_but_no_new_batch(self):
        self.client.force_login(self.engineer)
        response = self.client.get(reverse('inspections:dashboard'))
        self.assertContains(response, 'Inspections')
        # Engineer cannot manage -> "New Batch" link absent from nav.
        self.assertNotContains(response, reverse('inspections:batch_create'))

    def test_nav_hidden_for_outsider(self):
        # Outsider is forbidden from the dashboard entirely, so check a page they
        # CAN reach (the main dashboard) does not show the Inspections nav.
        self.client.force_login(self.outsider)
        response = self.client.get(reverse('dashboard:dashboard'))
        self.assertNotContains(response, reverse('inspections:dashboard'))
